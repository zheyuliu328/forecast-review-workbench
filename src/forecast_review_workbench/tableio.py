"""Bounded in-memory CSV/XLSX intake; never open a browser-supplied path."""

import base64
import binascii
import csv
import hashlib
import io
from datetime import date, datetime
from pathlib import PurePath
from zipfile import BadZipFile, ZipFile

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10_000
MAX_COLUMNS = 100


def _file_bytes(file):
    if not isinstance(file, dict) or not isinstance(file.get("name"), str) or not file["name"].strip():
        raise ValueError("Select a file with a nonempty name.")
    name = file["name"]
    if len(name) > 1024 or PurePath(name).suffix.lower() not in {".csv", ".xlsx"}:
        raise ValueError("Select a UTF-8 comma-separated .csv or a .xlsx file.")
    encoded = file.get("content_base64")
    if not isinstance(encoded, str) or len(encoded) > ((MAX_BYTES + 2) // 3) * 4:
        raise ValueError("Each original file must be at most 10 MiB and supplied as base64 bytes.")
    try:
        content = base64.b64decode(encoded, validate=True)
    except (binascii.Error, ValueError) as exc:
        raise ValueError("File bytes are not valid base64; select the original file again.") from exc
    if not content or len(content) > MAX_BYTES:
        raise ValueError("Select a nonempty file of at most 10 MiB.")
    return name, content, hashlib.sha256(content).hexdigest()


def _cell(value, kind=None):
    if kind is None:
        if value is None:
            kind = "blank"
        elif isinstance(value, bool):
            kind = "boolean"
        elif isinstance(value, (datetime, date)):
            kind = "date"
        elif isinstance(value, str):
            kind = "text"
        else:
            kind = "number"
    if kind == "formula" and not isinstance(value, str):
        formula_text = getattr(value, "text", None)
        text = formula_text if isinstance(formula_text, str) else f"[Excel {type(value).__name__} formula]"
    else:
        text = (
            "" if value is None else value.isoformat() if isinstance(value, (date, datetime)) else str(value)
        )
    return {"text": text, "kind": kind}


def _excel_rows(content, sheet, info):
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if (
                len(members) > 2048
                or len({item.filename for item in members}) != len(members)
                or sum(item.file_size for item in members) > 50 * 1024 * 1024
                or any(item.file_size > 20 * 1024 * 1024 or item.flag_bits & 1 for item in members)
            ):
                raise ValueError(
                    "Excel archive is encrypted, ambiguous or exceeds the expanded-size limit "
                    "(20 MiB/member, 50 MiB total)."
                )
            if "[Content_Types].xml" not in archive.namelist():
                raise ValueError(
                    "Excel workbook is missing [Content_Types].xml; select a valid exported .xlsx file."
                )
        from defusedxml.common import DefusedXmlException
        from defusedxml.ElementTree import ParseError
        from openpyxl import LXML, load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as exc:
        raise ValueError(
            "Excel support is unavailable; install the application's Excel dependencies."
        ) from exc
    except BadZipFile as exc:
        raise ValueError("Excel input is not a valid .xlsx ZIP archive.") from exc
    parser_errors = (
        KeyError,
        IndexError,
        TypeError,
        ValueError,
        AttributeError,
        OSError,
        BadZipFile,
        ParseError,
        DefusedXmlException,
        InvalidFileException,
    )
    if LXML:
        from lxml.etree import XMLSyntaxError

        parser_errors += (XMLSyntaxError,)
    workbook = None
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        info["sheets"] = list(workbook.sheetnames)
        if sheet is None:
            if len(info["sheets"]) != 1:
                raise ValueError("Select an Excel worksheet from the available sheet names.")
            sheet = info["sheets"][0]
        if sheet not in info["sheets"]:
            raise ValueError(f"Worksheet {sheet!r} was not found; choose one of the available sheets.")
        info["sheet"] = sheet
        worksheet = workbook[sheet]
        if not hasattr(worksheet, "reset_dimensions"):
            raise ValueError("Select a worksheet with cells; chart sheets are not tables.")
        worksheet.reset_dimensions()
        for row in worksheet.iter_rows():
            yield [
                _cell(
                    cell.value,
                    "formula" if cell.data_type == "f" else "error" if cell.data_type == "e" else None,
                )
                for cell in row
            ]
    except parser_errors as exc:
        raise ValueError(f"Cannot read Excel workbook structure: {exc}") from exc
    finally:
        if workbook is not None:
            workbook.close()


def _raw_rows(name, content, sheet, info):
    if PurePath(name).suffix.lower() == ".xlsx":
        yield from _excel_rows(content, sheet, info)
        return
    if sheet is not None:
        raise ValueError("CSV inputs have no worksheet; clear the sheet selection.")
    try:
        reader = csv.reader(io.StringIO(content.decode("utf-8-sig"), newline=""), strict=True)
        for row in reader:
            yield [_cell(value) for value in row]
    except (UnicodeError, csv.Error) as exc:
        raise ValueError(f"Cannot read UTF-8 comma-separated CSV: {exc}") from exc


def _info(sheet, header_row):
    return {
        "sheets": [],
        "sheet": sheet if isinstance(sheet, str) else None,
        "header_row": header_row if type(header_row) is int else 1,
        "headers": [],
        "preview": [],
        "row_count": 0,
        "error": None,
    }


def _read(file, sheet, header_row, info):
    if type(header_row) is not int or not 1 <= header_row <= 1000:
        raise ValueError("Header row must be an integer from 1 through 1000.")
    if sheet is not None and (not isinstance(sheet, str) or not sheet):
        raise ValueError("Choose a worksheet name or leave the sheet selection empty.")
    name, content, digest = _file_bytes(file)
    raw = []
    for row_number, cells in enumerate(_raw_rows(name, content, sheet, info), 1):
        if row_number > header_row + MAX_ROWS:
            raise ValueError("This table exceeds 10,000 physical data rows; select a smaller extract.")
        if len(cells) > MAX_COLUMNS:
            raise ValueError("This table exceeds 100 columns; select a smaller extract.")
        raw.append((row_number, cells))
        if len(info["preview"]) < 6:
            info["preview"].append([cell["text"] for cell in cells])
    if header_row > len(raw):
        raise ValueError(
            "The selected header row does not exist; inspect the raw preview and choose another row."
        )
    cells = list(raw[header_row - 1][1])
    excel = PurePath(name).suffix.lower() == ".xlsx"
    if excel:
        while cells and cells[-1]["kind"] == "blank":
            cells.pop()
    if not cells or any(cell["kind"] != "text" or not cell["text"].strip() for cell in cells):
        raise ValueError(
            "Header cells must be nonempty text; choose a row without formulas, blank or numeric headers."
        )
    headers = [cell["text"] for cell in cells]
    if len(set(headers)) != len(headers):
        raise ValueError("Duplicate headers are ambiguous; choose a table with distinct column names.")
    info["headers"] = headers
    rows, blanks = [], []
    for row_number, cells in raw[header_row:]:
        if not any(cell["text"] != "" for cell in cells):
            blanks.append(row_number)
            continue
        if len(cells) != len(headers):
            if not excel or any(cell["text"] for cell in cells[len(headers) :]):
                raise ValueError(
                    f"Row {row_number} has a different width from the header; "
                    "check the selected header and CSV delimiter."
                )
            cells = cells[: len(headers)] + [_cell(None)] * max(0, len(headers) - len(cells))
        rows.append({"row": row_number, "cells": dict(zip(headers, cells))})
    info["row_count"] = len(rows)
    info["preview"] = [[row["cells"][name]["text"] for name in headers] for row in rows[:6]]
    return {
        "file_name": name,
        "sha256": digest,
        "headers": headers,
        "rows": rows,
        "blank_rows_ignored": blanks,
        "sheet": info["sheet"],
        "header_row": header_row,
    }


def inspect_table(file, sheet=None, header_row=1):
    """Return selectable sheets, exact headers, a small preview and actionable errors."""
    info = _info(sheet, header_row)
    try:
        _read(file, sheet, header_row, info)
    except ValueError as exc:
        info["error"] = str(exc)
    return info


def read_table(file, sheet=None, header_row=1):
    """Internal engine intake; structural errors fail before any comparison."""
    return _read(file, sheet, header_row, _info(sheet, header_row))
